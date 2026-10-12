"""Derive a field-level Sankey summary from the local qiaopi workbook.

The raw workbook is never copied into the repository. Only aggregated
source-country / office / receiving-county counts are written.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def clean(value, fallback):
    if value is None:
        return fallback
    text = str(value).strip()
    return text or fallback


def normalize_source(value):
    value = clean(value, "未记录来源")
    return "印度尼西亚" if value == "印尼" else value


def top_known(counter, missing_label, n):
    return [key for key, _ in counter.most_common() if key != missing_label][:n]


def collapse(value, keep, missing_label, other_label):
    if value == missing_label:
        return missing_label
    return value if value in keep else other_label


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    index = {name: i for i, name in enumerate(header)}
    required = ["来源国", "批局名", "收批县"]
    missing_columns = [name for name in required if name not in index]
    if missing_columns:
        raise ValueError(f"missing columns: {missing_columns}")

    raw_source = Counter()
    raw_office = Counter()
    raw_county = Counter()
    records = []
    for row in rows:
        source = normalize_source(row[index["来源国"]])
        office = clean(row[index["批局名"]], "未记录批局")
        county = clean(row[index["收批县"]], "未记录收批县")
        raw_source[source] += 1
        raw_office[office] += 1
        raw_county[county] += 1
        records.append((source, office, county))

    source_keep = top_known(raw_source, "未记录来源", 6)
    office_keep = top_known(raw_office, "未记录批局", 10)
    county_keep = top_known(raw_county, "未记录收批县", 6)

    source_other = "其他来源"
    office_other = "其他批局"
    county_other = "其他收批县"
    source_nodes = source_keep + [source_other, "未记录来源"]
    office_nodes = office_keep + [office_other, "未记录批局"]
    county_nodes = county_keep + [county_other, "未记录收批县"]

    source_edges = Counter()
    office_edges = Counter()
    for source, office, county in records:
        source_view = collapse(source, source_keep, "未记录来源", source_other)
        office_view = collapse(office, office_keep, "未记录批局", office_other)
        county_view = collapse(county, county_keep, "未记录收批县", county_other)
        source_edges[(source_view, office_view)] += 1
        office_edges[(office_view, county_view)] += 1

    edges = []
    for (source, office), count in sorted(source_edges.items()):
        edges.append({"stage": "source-office", "source": source, "target": office, "count": count})
    for (office, county), count in sorted(office_edges.items()):
        edges.append({"stage": "office-county", "source": office, "target": county, "count": count})

    sample = next(
        (count for (source, office, county), count in Counter(records).items()
         if source == "新加坡" and office == "万益成" and county == "潮安"),
        0,
    )
    derived = {
        "schema_version": "qiaopi-panel-02-sankey-derived/v1",
        "source": "侨批可用数据(1).xlsx",
        "unit": "侨批记录条数",
        "rows": len(records),
        "columns_used": required,
        "normalization": {"source_aliases": {"印尼": "印度尼西亚"}},
        "display_rules": {
            "source_top_n": 6,
            "office_top_n": 10,
            "county_top_n": 6,
            "collapsed_labels": [source_other, office_other, county_other],
            "missing_labels": ["未记录来源", "未记录批局", "未记录收批县"],
        },
        "nodes": {"source": source_nodes, "office": office_nodes, "county": county_nodes},
        "raw_totals": {
            "source": raw_source,
            "office": raw_office,
            "county": raw_county,
        },
        "edges": edges,
        "sample_path": {"source": "新加坡", "office": "万益成", "county": "潮安", "count": sample},
    }

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "panel-02-sankey-derived-v1.json").write_text(
        json.dumps(derived, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    with (out_dir / "panel-02-sankey-edges-v1.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["stage", "source", "target", "count"])
        writer.writeheader()
        writer.writerows(edges)

    print(json.dumps({"rows": len(records), "sample_path_count": sample, "edges": len(edges)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
