"""Build an auditable office-reach summary from the local qiaopi workbook.

Only field-level aggregates are written to the repository; the raw workbook
is never copied or uploaded.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def clean(value, fallback):
    if value is None:
        return fallback
    text = str(value).strip()
    return text or fallback


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    index = {name: i for i, name in enumerate(header)}
    required = ["批局名", "收批县"]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"missing columns: {missing}")

    office_total = Counter()
    office_counties: dict[str, set[str]] = defaultdict(set)
    office_county_edges = Counter()
    rows_total = 0
    missing_office = 0
    missing_county = 0

    for row in rows:
        rows_total += 1
        office = clean(row[index["批局名"]], "未记录批局")
        county = clean(row[index["收批县"]], "未记录收批县")
        office_total[office] += 1
        if office == "未记录批局":
            missing_office += 1
        if county == "未记录收批县":
            missing_county += 1
        if office != "未记录批局" and county != "未记录收批县":
            office_counties[office].add(county)
            office_county_edges[(office, county)] += 1

    named = []
    for office, total in office_total.most_common():
        if office == "未记录批局":
            continue
        named.append(
            {
                "office": office,
                "records": total,
                "known_county_count": len(office_counties.get(office, set())),
            }
        )

    derived = {
        "schema_version": "qiaopi-office-reach-derived/v1",
        "source": "侨批可用数据(1).xlsx",
        "unit": {"records": "侨批记录条数", "known_county_count": "已知收批县类别数（个）"},
        "rows": rows_total,
        "columns_used": required,
        "missing": {
            "office_rows": missing_office,
            "county_rows": missing_county,
            "office_label": "未记录批局",
            "county_label": "未记录收批县",
        },
        "offices": named,
        "top_offices": named[:12],
        "office_county_edges": [
            {"office": office, "county": county, "count": count}
            for (office, county), count in sorted(office_county_edges.items())
        ],
    }
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(derived, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"rows": rows_total, "offices": len(named), "missing_office": missing_office}, ensure_ascii=False))


if __name__ == "__main__":
    main()

